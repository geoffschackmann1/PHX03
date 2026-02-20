"""
Cap Dashboard — Excel Report Builder
Generates a multi-tab Excel workbook with cap position data.
Mirrors the QPi excel_builder.py styling pattern.
"""
import logging
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from cap_config import (
    AGENCY_NAME,
    CAP_YEARS,
    RISK_TIERS,
    CURRENCY_FORMAT,
    PCT_FORMAT,
)

logger = logging.getLogger(__name__)

# =============================================================================
# STYLE DEFINITIONS (mirrors QPi pattern)
# =============================================================================
HF = Font(name="Arial", bold=True, size=10, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F4E79")
TF = Font(name="Arial", bold=True, size=14, color="1F4E79")
SF = Font(name="Arial", bold=True, size=12, color="1F4E79")
BF = Font(name="Arial", bold=True, size=10)
DF = Font(name="Arial", size=10)
GF = Font(name="Arial", bold=True, size=10, color="006100")
RF = Font(name="Arial", bold=True, size=10, color="9C0006")
AF = Font(name="Arial", bold=True, size=10, color="9C6500")
GFL = PatternFill("solid", fgColor="C6EFCE")
RFL = PatternFill("solid", fgColor="FFC7CE")
AFL = PatternFill("solid", fgColor="FFEB9C")
HIGH_FILL = PatternFill("solid", fgColor="FFC7CE")
WATCH_FILL = PatternFill("solid", fgColor="FFEB9C")
HEALTHY_FILL = PatternFill("solid", fgColor="C6EFCE")
TB = Border(
    left=Side("thin"), right=Side("thin"),
    top=Side("thin"), bottom=Side("thin"),
)
CA = Alignment(horizontal="center", vertical="center", wrap_text=True)
LA = Alignment(horizontal="left", vertical="center", wrap_text=True)
RA = Alignment(horizontal="right", vertical="center", wrap_text=True)


def _header_row(ws, row: int, headers: list[str]):
    """Write a styled header row."""
    for c, h in enumerate(headers, 1):
        cl = ws.cell(row=row, column=c, value=h)
        cl.font = HF
        cl.fill = HFILL
        cl.alignment = CA
        cl.border = TB


def _cell(ws, r: int, c: int, val=None, fmt: str = None):
    """Write a styled data cell."""
    cl = ws.cell(row=r, column=c)
    if val is not None:
        cl.value = val
    cl.font = DF
    cl.alignment = CA
    cl.border = TB
    if fmt:
        cl.number_format = fmt
    return cl


def _risk_fill(tier: str) -> PatternFill:
    if tier == "HIGH":
        return HIGH_FILL
    elif tier == "WATCH":
        return WATCH_FILL
    return HEALTHY_FILL


# =============================================================================
# WORKBOOK BUILDER
# =============================================================================

def build_cap_workbook(
    summary: dict,
    patients_df: pd.DataFrame,
    payments_df: pd.DataFrame,
    monthly_df: pd.DataFrame,
    output_dir: str = "./output",
    cap_year_key: str = "FY2025",
) -> str:
    """
    Build the cap dashboard Excel workbook.

    Tabs:
    1. Cap Summary — high-level position metrics
    2. Patient Detail — per-patient cap credit and risk tier
    3. Payments — claim payment detail
    4. Monthly Trend — monthly reimbursement trend

    Returns the output file path.
    """
    wb = openpyxl.Workbook()

    # --- Tab 1: Cap Summary ---
    _build_summary_tab(wb, summary)

    # --- Tab 2: Patient Detail ---
    _build_patient_tab(wb, patients_df)

    # --- Tab 3: Payments ---
    _build_payments_tab(wb, payments_df)

    # --- Tab 4: Monthly Trend ---
    _build_monthly_tab(wb, monthly_df, summary)

    # Remove default sheet if extras were created
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]

    # Save
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    filename = f"Cap_Dashboard_{cap_year_key}_{date.today().isoformat()}.xlsx"
    filepath = output_path / filename
    wb.save(str(filepath))
    logger.info(f"Cap workbook saved: {filepath}")
    return str(filepath)


# =============================================================================
# TAB BUILDERS
# =============================================================================

def _build_summary_tab(wb: openpyxl.Workbook, summary: dict):
    """Tab 1: Cap Position Summary."""
    ws = wb.active
    ws.title = "Cap Summary"
    ws.sheet_properties.tabColor = "1F4E79"

    # Title
    ws.merge_cells("A1:D1")
    cl = ws.cell(row=1, column=1, value=f"{AGENCY_NAME} — Medicare Hospice Cap Dashboard")
    cl.font = TF
    cl.alignment = LA

    ws.merge_cells("A2:D2")
    cl = ws.cell(row=2, column=1, value=summary.get("cap_year_label", summary["cap_year"]))
    cl.font = SF
    cl.alignment = LA

    ws.merge_cells("A3:D3")
    cl = ws.cell(row=3, column=1, value=f"Generated: {datetime.now().strftime('%m/%d/%Y %I:%M %p')}")
    cl.font = DF
    cl.alignment = LA

    # --- Kinnser Estimate Section ---
    row = 5
    ws.merge_cells(f"A{row}:D{row}")
    cl = ws.cell(row=row, column=1, value="Kinnser/WellSky Estimate (from Snowflake)")
    cl.font = SF

    metrics = [
        ("Cap Rate (per beneficiary)", summary["cap_rate"], "$#,##0.00"),
        ("Est. Pro-Rated Beneficiaries", summary["est_bene_count"], "#,##0.0000"),
        ("Est. Cap Allowable", summary["est_allowable"], "$#,##0.00"),
        ("Est. Net Reimbursement", summary.get("est_net_pymts"), "$#,##0.00"),
        ("Est. Cap Position (+ = Under Cap)", summary.get("est_position"), "$#,##0.00"),
        ("Est. Cap Utilization %", summary.get("est_utilization_pct"), "0.0%"),
    ]

    row += 1
    _header_row(ws, row, ["Metric", "Value", "", ""])
    for label, val, fmt in metrics:
        row += 1
        _cell(ws, row, 1, label).alignment = LA
        _cell(ws, row, 1).font = BF
        c = _cell(ws, row, 2, val, fmt)
        if label == "Est. Cap Position (+ = Under Cap)" and val is not None:
            c.font = GF if val >= 0 else RF
            c.fill = GFL if val >= 0 else RFL

    # --- PS&R Authoritative Section ---
    row += 2
    ws.merge_cells(f"A{row}:D{row}")
    cl = ws.cell(row=row, column=1, value="PS&R Authoritative (from CMS Filing)")
    cl.font = SF

    psr_metrics = [
        ("PS&R Pro-Rated Beneficiaries", summary.get("auth_bene_count"), "#,##0.0000"),
        ("PS&R Allowable", summary.get("auth_allowable"), "$#,##0.00"),
        ("PS&R Gross Reimbursement", summary.get("auth_gross_pymts"), "$#,##0.00"),
        ("PS&R Sequestration", summary.get("auth_sequestration"), "$#,##0.00"),
        ("PS&R Net Reimbursement", summary.get("auth_net_pymts"), "$#,##0.00"),
        ("PS&R Cap Position (+ = Under Cap)", summary.get("auth_position"), "$#,##0.00"),
        ("PS&R Utilization %", summary.get("auth_utilization_pct"), "0.0%"),
    ]

    row += 1
    _header_row(ws, row, ["Metric", "Value", "", ""])
    for label, val, fmt in psr_metrics:
        row += 1
        _cell(ws, row, 1, label).alignment = LA
        _cell(ws, row, 1).font = BF
        if val is not None:
            c = _cell(ws, row, 2, val, fmt)
            if "Position" in label:
                c.font = GF if val >= 0 else RF
                c.fill = GFL if val >= 0 else RFL
        else:
            _cell(ws, row, 2, "Not yet filed")

    # --- Census & Risk ---
    row += 2
    ws.merge_cells(f"A{row}:D{row}")
    cl = ws.cell(row=row, column=1, value="Census & Risk Distribution")
    cl.font = SF

    risk_metrics = [
        ("Total Medicare Patients (Cap Year)", summary["total_patients"]),
        ("Active Census (Medicare)", summary["active_census"]),
        ("High Risk (BP >= 5)", summary.get("high_risk_count", 0)),
        ("Watch (BP 3-4)", summary.get("watch_count", 0)),
        ("Healthy (BP 1-2)", summary.get("healthy_count", 0)),
        ("Inherited (Pre-Acquisition)", summary.get("inherited_count", 0)),
    ]

    row += 1
    _header_row(ws, row, ["Metric", "Count", "", ""])
    for label, val in risk_metrics:
        row += 1
        _cell(ws, row, 1, label).alignment = LA
        _cell(ws, row, 1).font = BF
        c = _cell(ws, row, 2, val, "#,##0")
        if "High Risk" in label and val > 0:
            c.fill = HIGH_FILL
            c.font = RF
        elif "Watch" in label and val > 0:
            c.fill = WATCH_FILL
            c.font = AF

    # Column widths
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 5
    ws.column_dimensions["D"].width = 5


def _build_patient_tab(wb: openpyxl.Workbook, patients_df: pd.DataFrame):
    """Tab 2: Patient-level cap detail."""
    ws = wb.create_sheet("Patient Detail")
    ws.sheet_properties.tabColor = "4472C4"

    if patients_df.empty:
        ws.cell(row=1, column=1, value="No patient data available.")
        return

    headers = [
        "Last Name", "First Name", "MRN", "Medicare #",
        "SOC Date", "Discharge Date", "Max BP#",
        "Risk Tier", "LOC Days", "Cap Credit",
        "Net Reimb.", "Active", "Inherited",
    ]
    _header_row(ws, 1, headers)

    cols_map = [
        ("last_name", None),
        ("first_name", None),
        ("medical_record_number", None),
        ("medicare_number", None),
        ("start_of_care_date", "MM/DD/YYYY"),
        ("discharge_date", "MM/DD/YYYY"),
        ("max_bp", "#,##0"),
        ("risk_tier", None),
        ("total_loc_days", "#,##0"),
        ("cap_credit", "0.0000"),
        ("total_net_reimbursement", "$#,##0.00"),
        ("is_active", None),
        ("is_inherited", None),
    ]

    # Sort by risk tier (HIGH first), then by cap credit desc
    tier_order = {"HIGH": 0, "WATCH": 1, "HEALTHY": 2}
    df = patients_df.copy()
    df["_tier_order"] = df["risk_tier"].map(tier_order).fillna(3)
    df = df.sort_values(["_tier_order", "cap_credit"], ascending=[True, False])

    for i, (_, row_data) in enumerate(df.iterrows(), start=2):
        for j, (col, fmt) in enumerate(cols_map, start=1):
            val = row_data.get(col)
            if isinstance(val, bool):
                val = "Yes" if val else ""
            elif hasattr(val, "strftime") and fmt == "MM/DD/YYYY":
                val = val.strftime("%m/%d/%Y") if pd.notna(val) else ""
                fmt = None
            c = _cell(ws, i, j, val, fmt)
            if col == "risk_tier":
                c.fill = _risk_fill(str(val))

    # Column widths
    widths = [14, 12, 12, 14, 12, 12, 8, 10, 10, 10, 14, 8, 10]
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    # Freeze top row
    ws.freeze_panes = "A2"


def _build_payments_tab(wb: openpyxl.Workbook, payments_df: pd.DataFrame):
    """Tab 3: Claim payment detail."""
    ws = wb.create_sheet("Payments")
    ws.sheet_properties.tabColor = "70AD47"

    if payments_df.empty:
        ws.cell(row=1, column=1, value="No payment data available.")
        return

    headers = [
        "Last Name", "First Name", "MRN", "Claim #",
        "Claim Start", "Claim End", "Check Date",
        "Payment", "Adjustment", "Net Reimb.",
    ]
    _header_row(ws, 1, headers)

    cols = [
        ("patient_last_name", None),
        ("patient_first_name", None),
        ("medical_record_number", None),
        ("claim_number", "#,##0"),
        ("claim_start_date", "MM/DD/YYYY"),
        ("claim_end_date", "MM/DD/YYYY"),
        ("check_date", "MM/DD/YYYY"),
        ("claim_payment_applied", "$#,##0.00"),
        ("claim_adjustment", "$#,##0.00"),
        ("claim_net_reimbursement", "$#,##0.00"),
    ]

    df = payments_df.sort_values("check_date", ascending=False) if "check_date" in payments_df.columns else payments_df

    for i, (_, row_data) in enumerate(df.iterrows(), start=2):
        for j, (col, fmt) in enumerate(cols, start=1):
            val = row_data.get(col)
            if hasattr(val, "strftime") and fmt == "MM/DD/YYYY":
                val = val.strftime("%m/%d/%Y") if pd.notna(val) else ""
                fmt = None
            _cell(ws, i, j, val, fmt)

    # Totals row
    total_row = len(df) + 2
    _cell(ws, total_row, 1, "TOTAL").font = BF
    for j, (col, fmt) in enumerate(cols, 1):
        if col in ("claim_payment_applied", "claim_adjustment", "claim_net_reimbursement"):
            val = pd.to_numeric(df[col], errors="coerce").sum()
            c = _cell(ws, total_row, j, val, "$#,##0.00")
            c.font = BF

    widths = [14, 12, 12, 10, 12, 12, 12, 14, 14, 14]
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    ws.freeze_panes = "A2"


def _build_monthly_tab(
    wb: openpyxl.Workbook,
    monthly_df: pd.DataFrame,
    summary: dict,
):
    """Tab 4: Monthly reimbursement trend."""
    ws = wb.create_sheet("Monthly Trend")
    ws.sheet_properties.tabColor = "ED7D31"

    if monthly_df.empty:
        ws.cell(row=1, column=1, value="No monthly trend data available.")
        return

    headers = [
        "Month", "Net Reimbursement", "Claim Count",
        "Cumulative Reimb.", "Cap Allowable", "Utilization %",
    ]
    _header_row(ws, 1, headers)

    cum_reimb = 0
    allowable = summary.get("est_allowable", 0)

    for i, (_, row_data) in enumerate(monthly_df.iterrows(), start=2):
        month_val = row_data.get("month")
        if hasattr(month_val, "strftime"):
            month_val = month_val.strftime("%b %Y")
        net = float(row_data.get("total_net_reimbursement", 0))
        cum_reimb += net
        util_pct = (cum_reimb / allowable * 100) if allowable else 0

        _cell(ws, i, 1, month_val)
        _cell(ws, i, 2, net, "$#,##0.00")
        _cell(ws, i, 3, int(row_data.get("claim_count", 0)), "#,##0")
        _cell(ws, i, 4, cum_reimb, "$#,##0.00")
        _cell(ws, i, 5, allowable, "$#,##0.00")
        c = _cell(ws, i, 6, util_pct / 100, "0.0%")
        if util_pct >= 95:
            c.fill = HIGH_FILL
            c.font = RF
        elif util_pct >= 85:
            c.fill = WATCH_FILL
            c.font = AF

    widths = [12, 18, 12, 18, 18, 14]
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    ws.freeze_panes = "A2"
