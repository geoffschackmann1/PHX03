"""
QPi Scorecard — Mock Pipeline Validation Test
Runs the full pipeline with realistic mock data (no Snowflake needed).
Validates: scorecard_engine, excel_builder, config, and main output flow.
"""
import logging
import sys
from datetime import date
from pathlib import Path

import pandas as pd

# Load env so config module picks up variables
from load_env import load_dotenv
load_dotenv()

from config import get_pay_period, AGENCIES
from scorecard_engine import build_clinician_scorecard, build_agency_metrics, evaluate_target
from excel_builder import build_workbook

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger("test_mock")


# =============================================================================
# MOCK DATA — matches Snowflake query return schemas exactly
# =============================================================================

def mock_productivity() -> pd.DataFrame:
    """Mock output of SQL_CLINICIAN_PRODUCTIVITY (15 clinicians)."""
    data = [
        # clinician_name, discipline, emp_type, skill_level, soc_eval_ct, recert_ct, routine_ct, discharge_ct, total_visits, missed_visits, unique_patients, pct_soc_eval, pct_routine
        ("DAVID BERRY", "PT", "Full Time", "SKILLED", 8, 3, 10, 3, 24, 1, 12, 33.3, 41.7),
        ("ARNOLD GONZALES", "OT", "Full Time", "SKILLED", 5, 2, 12, 2, 21, 0, 10, 23.8, 57.1),
        ("MARIA SANTOS", "SN", "Full Time", "SKILLED", 6, 1, 14, 3, 24, 2, 15, 25.0, 58.3),
        ("JENNIFER WRIGHT", "PT", "Full Time", "SKILLED", 4, 2, 8, 1, 15, 0, 8, 26.7, 53.3),
        ("MICHAEL CHEN", "OT", "Full Time", "SKILLED", 3, 1, 10, 2, 16, 1, 9, 18.8, 62.5),
        ("SARAH JOHNSON", "SN", "Full Time", "SKILLED", 7, 2, 6, 2, 17, 0, 11, 41.2, 35.3),
        ("ROBERT WILLIAMS", "ST", "Full Time", "SKILLED", 2, 1, 8, 1, 12, 0, 7, 16.7, 66.7),
        ("LISA MARTINEZ", "MSW", "Full Time", "SKILLED", 1, 0, 6, 1, 8, 0, 6, 12.5, 75.0),
        ("JAMES THOMPSON", "PT", "PRN", "SKILLED", 2, 0, 5, 0, 7, 0, 4, 28.6, 71.4),
        ("PATRICIA GARCIA", "PTA", "Full Time", "ASSISTANT", 0, 0, 12, 2, 14, 0, 8, 0.0, 85.7),
        ("DANIEL BROWN", "COTA", "Full Time", "ASSISTANT", 0, 0, 10, 1, 11, 1, 7, 0.0, 90.9),
        ("NANCY LEE", "LPN/LVN", "Full Time", "ASSISTANT", 0, 0, 8, 0, 8, 0, 5, 0.0, 100.0),
        ("THOMAS DAVIS", "HHA", "Part Time", "ASSISTANT", 0, 0, 6, 0, 6, 0, 4, 0.0, 100.0),
        ("AMANDA WILSON", "SN", "PRN", "SKILLED", 1, 0, 3, 1, 5, 0, 3, 20.0, 60.0),
        ("KEVIN MOORE", "PT", "Full Time", "SKILLED", 3, 1, 7, 1, 12, 0, 6, 25.0, 58.3),
    ]
    return pd.DataFrame(data, columns=[
        "clinician_name", "discipline", "emp_type", "skill_level",
        "soc_eval_ct", "recert_ct", "routine_ct", "discharge_ct",
        "total_visits", "missed_visits", "unique_patients",
        "pct_soc_eval", "pct_routine",
    ])


def mock_documentation() -> pd.DataFrame:
    """Mock output of SQL_DOCUMENTATION."""
    data = [
        ("DAVID BERRY", "PT", 24, 22, 2, 91.7, 1.2, 0, 4),
        ("ARNOLD GONZALES", "OT", 21, 19, 2, 90.5, 1.5, 0, 5),
        ("MARIA SANTOS", "SN", 24, 20, 4, 83.3, 1.8, 0, 6),
        ("JENNIFER WRIGHT", "PT", 15, 14, 1, 93.3, 0.9, 0, 3),
        ("MICHAEL CHEN", "OT", 16, 13, 3, 81.3, 2.1, 0, 7),
        ("SARAH JOHNSON", "SN", 17, 16, 1, 94.1, 0.8, 0, 2),
        ("ROBERT WILLIAMS", "ST", 12, 11, 1, 91.7, 1.1, 0, 3),
        ("LISA MARTINEZ", "MSW", 8, 8, 0, 100.0, 0.5, 0, 1),
        ("JAMES THOMPSON", "PT", 7, 6, 1, 85.7, 1.4, 0, 4),
        ("PATRICIA GARCIA", "PTA", 14, 13, 1, 92.9, 1.0, 0, 3),
        ("DANIEL BROWN", "COTA", 11, 9, 2, 81.8, 2.0, 0, 5),
        ("NANCY LEE", "LPN/LVN", 8, 7, 1, 87.5, 1.3, 0, 4),
        ("THOMAS DAVIS", "HHA", 6, 6, 0, 100.0, 0.3, 0, 1),
        ("AMANDA WILSON", "SN", 5, 4, 1, 80.0, 2.5, 1, 8),
        ("KEVIN MOORE", "PT", 12, 11, 1, 91.7, 1.2, 0, 3),
    ]
    return pd.DataFrame(data, columns=[
        "clinician_name", "discipline", "total_docs", "docs_on_time",
        "docs_late", "on_time_pct", "avg_days_to_submit",
        "fastest_days", "slowest_days",
    ])


def mock_payroll() -> pd.DataFrame:
    """Mock iSolved payroll data matching the adapter's PAYROLL_SCHEMA."""
    data = [
        ("DAVID BERRY", 80, 4, 4230.77, 125.50, 4356.27, 0, 0, 0, 0, 0, 0),
        ("ARNOLD GONZALES", 80, 0, 4423.08, 98.00, 4521.08, 0, 0, 0, 0, 0, 0),
        ("MARIA SANTOS", 80, 6, 3846.15, 145.00, 3991.15, 0, 8, 0, 0, 0, 8),
        ("JENNIFER WRIGHT", 72, 0, 3600.00, 85.00, 3685.00, 8, 0, 0, 0, 0, 8),
        ("MICHAEL CHEN", 80, 0, 4000.00, 110.00, 4110.00, 0, 0, 0, 0, 0, 0),
        ("SARAH JOHNSON", 80, 2, 4100.00, 130.00, 4230.00, 0, 0, 0, 0, 0, 0),
        ("ROBERT WILLIAMS", 80, 0, 3500.00, 75.00, 3575.00, 0, 0, 0, 0, 0, 0),
        ("LISA MARTINEZ", 80, 0, 3200.00, 60.00, 3260.00, 0, 0, 0, 0, 0, 0),
        ("JAMES THOMPSON", 40, 0, 1800.00, 50.00, 1850.00, 0, 0, 0, 0, 0, 0),
        ("PATRICIA GARCIA", 80, 0, 2800.00, 140.00, 2940.00, 0, 0, 0, 0, 0, 0),
        ("DANIEL BROWN", 80, 0, 2600.00, 95.00, 2695.00, 0, 0, 0, 0, 0, 0),
        ("NANCY LEE", 80, 0, 2400.00, 80.00, 2480.00, 0, 0, 0, 0, 0, 0),
        ("THOMAS DAVIS", 40, 0, 1200.00, 45.00, 1245.00, 0, 0, 0, 0, 0, 0),
        ("AMANDA WILSON", 24, 0, 1080.00, 35.00, 1115.00, 0, 0, 0, 0, 0, 0),
        ("KEVIN MOORE", 80, 0, 4000.00, 100.00, 4100.00, 0, 0, 0, 0, 0, 0),
    ]
    return pd.DataFrame(data, columns=[
        "clinician_name", "regular_hours", "overtime_hours", "gross_wages",
        "mileage", "total_cost", "vacation_hours", "pto_hours",
        "sick_hours", "holiday_hours", "bereavement_hours", "time_off_hours",
    ])


def mock_snowflake_data() -> dict:
    """Mock output of SnowflakeClient.get_all() — all 8 queries."""
    return {
        "productivity": mock_productivity(),
        "documentation": mock_documentation(),
        "census": pd.DataFrame([{"census_count": 44}]),
        "lupa": pd.DataFrame([{"total_episodes": 85, "lupa_episodes": 5, "lupa_pct": 5.9}]),
        "non_admits": pd.DataFrame([{"non_admit_count": 2}]),
        "visits_per_episode": pd.DataFrame([{
            "total_episodes": 18, "total_visits": 142, "visits_per_episode": 7.9,
        }]),
        "soc_medicare": pd.DataFrame([{
            "total_socs": 42, "medicare_socs": 39, "medicare_pct": 92.9,
        }]),
        "assistant_util": pd.DataFrame([
            {"skill_level": "ASSISTANT", "visit_count": 39},
            {"skill_level": "SKILLED", "visit_count": 161},
        ]),
    }


def mock_prior_productivity() -> pd.DataFrame:
    """Mock prior period productivity (slightly different numbers for trending)."""
    data = [
        ("DAVID BERRY", "PT", "Full Time", "SKILLED", 7, 2, 9, 2, 20, 0, 11, 35.0, 45.0),
        ("ARNOLD GONZALES", "OT", "Full Time", "SKILLED", 4, 1, 11, 2, 18, 0, 9, 22.2, 61.1),
        ("MARIA SANTOS", "SN", "Full Time", "SKILLED", 5, 1, 12, 2, 20, 1, 13, 25.0, 60.0),
        ("SARAH JOHNSON", "SN", "Full Time", "SKILLED", 6, 2, 5, 1, 14, 0, 9, 42.9, 35.7),
        ("KEVIN MOORE", "PT", "Full Time", "SKILLED", 2, 1, 6, 1, 10, 0, 5, 20.0, 60.0),
    ]
    return pd.DataFrame(data, columns=[
        "clinician_name", "discipline", "emp_type", "skill_level",
        "soc_eval_ct", "recert_ct", "routine_ct", "discharge_ct",
        "total_visits", "missed_visits", "unique_patients",
        "pct_soc_eval", "pct_routine",
    ])


def mock_prior_documentation() -> pd.DataFrame:
    """Mock prior period documentation."""
    data = [
        ("DAVID BERRY", "PT", 20, 18, 2, 90.0, 1.3, 0, 4),
        ("ARNOLD GONZALES", "OT", 18, 16, 2, 88.9, 1.6, 0, 5),
        ("MARIA SANTOS", "SN", 20, 16, 4, 80.0, 2.0, 0, 7),
        ("SARAH JOHNSON", "SN", 14, 13, 1, 92.9, 0.9, 0, 3),
        ("KEVIN MOORE", "PT", 10, 9, 1, 90.0, 1.2, 0, 3),
    ]
    return pd.DataFrame(data, columns=[
        "clinician_name", "discipline", "total_docs", "docs_on_time",
        "docs_late", "on_time_pct", "avg_days_to_submit",
        "fastest_days", "slowest_days",
    ])


# =============================================================================
# TEST EXECUTION
# =============================================================================

def run_test():
    pp = get_pay_period(4)  # PP4: 1/25/2026 - 2/7/2026
    clinic_key = 3743
    agency = AGENCIES[clinic_key]
    errors = []

    print("=" * 60)
    print("QPi MOCK PIPELINE VALIDATION")
    print(f"Pay Period: {pp['label']} ({pp['range_str']})")
    print(f"Agency: {agency.name} ({clinic_key})")
    print("=" * 60)

    # ------------------------------------------------------------------
    # Test 1: Pay period calculation
    # ------------------------------------------------------------------
    print("\n[TEST 1] Pay period calculation...")
    assert pp["start"] == date(2026, 1, 25), f"PP4 start wrong: {pp['start']}"
    assert pp["end"] == date(2026, 2, 7), f"PP4 end wrong: {pp['end']}"
    assert pp["number"] == 4, f"PP4 number wrong: {pp['number']}"
    print(f"  PASS: PP4 = {pp['start']} to {pp['end']}")

    # ------------------------------------------------------------------
    # Test 2: Scorecard engine — no payroll
    # ------------------------------------------------------------------
    print("\n[TEST 2] Scorecard engine (Snowflake only, no payroll)...")
    sf_data = mock_snowflake_data()
    df_no_pay = build_clinician_scorecard(
        sf_productivity=sf_data["productivity"],
        sf_documentation=sf_data["documentation"],
        payroll=pd.DataFrame(),
    )
    assert len(df_no_pay) == 15, f"Expected 15 clinicians, got {len(df_no_pay)}"
    assert "points" in df_no_pay.columns, "Missing 'points' column"
    assert "adj_target" in df_no_pay.columns, "Missing 'adj_target' column"
    assert "meets_target" in df_no_pay.columns, "Missing 'meets_target' column"
    assert df_no_pay["total_visits"].sum() == 200, f"Total visits: {df_no_pay['total_visits'].sum()}"

    # Points check: DAVID BERRY: 8*2.5 + 3*1.5 + 10*1.0 + 3*1.0 = 20+4.5+10+3 = 37.5
    berry = df_no_pay[df_no_pay["clinician_name"] == "DAVID BERRY"].iloc[0]
    assert berry["points"] == 37.5, f"Berry points wrong: {berry['points']}"
    print(f"  PASS: 15 clinicians, 200 visits, Berry=37.5pts")

    # Utilization should be NA without payroll
    assert pd.isna(df_no_pay["utilization_pct"].iloc[0]) or df_no_pay["regular_hours"].iloc[0] == 0, \
        "Utilization should be NA without payroll"
    print(f"  PASS: Utilization correctly NA without payroll")

    # ------------------------------------------------------------------
    # Test 3: Scorecard engine — with payroll
    # ------------------------------------------------------------------
    print("\n[TEST 3] Scorecard engine (Snowflake + iSolved)...")
    payroll = mock_payroll()
    df = build_clinician_scorecard(
        sf_productivity=sf_data["productivity"],
        sf_documentation=sf_data["documentation"],
        payroll=payroll,
    )
    assert len(df) == 15, f"Expected 15 clinicians, got {len(df)}"

    berry = df[df["clinician_name"] == "DAVID BERRY"].iloc[0]
    assert berry["regular_hours"] == 80, f"Berry regular_hours wrong: {berry['regular_hours']}"
    assert berry["total_cost"] == 4356.27, f"Berry total_cost wrong: {berry['total_cost']}"

    # CPV: 4356.27 / 24 = 181.51
    assert abs(berry["cpv"] - 181.51) < 0.1, f"Berry CPV wrong: {berry['cpv']}"

    # Utilization: regular_hours / bp_hours * 100
    # BP hours: 8*2.75 + 3*1.75 + 10*1.25 + 3*1.25 = 22+5.25+12.5+3.75 = 43.5
    assert abs(berry["bp_hours"] - 43.5) < 0.1, f"Berry BP hours wrong: {berry['bp_hours']}"
    expected_util = round(80 / 43.5 * 100, 1)
    assert abs(berry["utilization_pct"] - expected_util) < 0.5, \
        f"Berry utilization wrong: {berry['utilization_pct']} (expected ~{expected_util})"

    # Points per hour: 37.5 / 80 = 0.47
    expected_pph = round(37.5 / 80, 2)
    assert abs(berry["points_per_hour"] - expected_pph) < 0.01, \
        f"Berry pts/hr wrong: {berry['points_per_hour']} (expected {expected_pph})"

    # Adjusted target for JENNIFER WRIGHT (8 hrs time off)
    wright = df[df["clinician_name"] == "JENNIFER WRIGHT"].iloc[0]
    expected_adj = round(30.0 * (80 - 8) / 80, 1)
    assert abs(wright["adj_target"] - expected_adj) < 0.1, \
        f"Wright adj_target wrong: {wright['adj_target']} (expected {expected_adj})"
    print(f"  PASS: Payroll merged, CPV/Util/PPH calculated correctly")

    # ------------------------------------------------------------------
    # Test 4: Documentation merge
    # ------------------------------------------------------------------
    print("\n[TEST 4] Documentation merge...")
    assert "on_time_pct" in df.columns, "Missing on_time_pct"
    assert "docs_late" in df.columns, "Missing docs_late"
    berry_doc = df[df["clinician_name"] == "DAVID BERRY"].iloc[0]
    assert berry_doc["on_time_pct"] == 91.7, f"Berry on_time_pct wrong: {berry_doc['on_time_pct']}"
    assert berry_doc["docs_late"] == 2, f"Berry docs_late wrong: {berry_doc['docs_late']}"
    print(f"  PASS: Doc metrics merged correctly")

    # ------------------------------------------------------------------
    # Test 5: Prior period trending
    # ------------------------------------------------------------------
    print("\n[TEST 5] Prior period trending...")
    df_trend = build_clinician_scorecard(
        sf_productivity=sf_data["productivity"],
        sf_documentation=sf_data["documentation"],
        payroll=payroll,
        prior_productivity=mock_prior_productivity(),
        prior_documentation=mock_prior_documentation(),
    )
    berry_trend = df_trend[df_trend["clinician_name"] == "DAVID BERRY"].iloc[0]
    # Prior Berry: 7*2.5 + 2*1.5 + 9*1.0 + 2*1.0 = 17.5+3+9+2 = 31.5
    assert abs(berry_trend["prior_points"] - 31.5) < 0.1, \
        f"Berry prior_points wrong: {berry_trend['prior_points']}"
    assert abs(berry_trend["points_trend"] - (37.5 - 31.5)) < 0.1, \
        f"Berry points_trend wrong: {berry_trend['points_trend']}"
    assert abs(berry_trend["prior_on_time_pct"] - 90.0) < 0.1, \
        f"Berry prior_on_time_pct wrong: {berry_trend['prior_on_time_pct']}"

    # Clinicians without prior data should have NA
    wilson = df_trend[df_trend["clinician_name"] == "AMANDA WILSON"].iloc[0]
    assert pd.isna(wilson["prior_points"]), "Wilson should have NA prior_points"
    print(f"  PASS: Trending calculated correctly, NA for missing priors")

    # ------------------------------------------------------------------
    # Test 6: Agency metrics
    # ------------------------------------------------------------------
    print("\n[TEST 6] Agency metrics...")
    am = build_agency_metrics(sf_data)
    assert am.census == 44, f"Census wrong: {am.census}"
    assert am.lupa_pct == 5.9, f"LUPA % wrong: {am.lupa_pct}"
    assert am.non_admits == 2, f"Non-admits wrong: {am.non_admits}"
    assert am.visits_per_episode == 7.9, f"Visits/episode wrong: {am.visits_per_episode}"
    assert am.soc_medicare_pct == 92.9, f"SoC Medicare % wrong: {am.soc_medicare_pct}"
    assert am.assistant_visits == 39, f"Assistant visits wrong: {am.assistant_visits}"
    assert am.skilled_visits == 161, f"Skilled visits wrong: {am.skilled_visits}"
    expected_asst_pct = round(39 / (39 + 161) * 100, 1)
    assert abs(am.assistant_pct - expected_asst_pct) < 0.1, \
        f"Assistant % wrong: {am.assistant_pct} (expected {expected_asst_pct})"
    print(f"  PASS: All agency metrics correct")

    # ------------------------------------------------------------------
    # Test 7: evaluate_target
    # ------------------------------------------------------------------
    print("\n[TEST 7] Target evaluation...")
    assert evaluate_target(95.0, "doc_on_time_pct") == "MEETS"
    assert evaluate_target(85.0, "doc_on_time_pct") == "BELOW"
    assert evaluate_target(100.0, "cpv") == "MEETS"
    assert evaluate_target(110.0, "cpv") == "ABOVE"
    assert evaluate_target(8.5, "visits_per_episode") == "MEETS"
    assert evaluate_target(9.0, "visits_per_episode") == "ABOVE"
    assert evaluate_target(None, "census") == "N/A"
    assert evaluate_target(130, "census") == "MEETS"
    assert evaluate_target(100, "census") == "BELOW"
    print(f"  PASS: All target evaluations correct")

    # ------------------------------------------------------------------
    # Test 8: OPP eligibility
    # ------------------------------------------------------------------
    print("\n[TEST 8] OPP eligibility...")
    opp_eligible = df[df["opp_eligible"] == True]
    opp_not = df[df["opp_eligible"] == False]
    print(f"  OPP eligible: {len(opp_eligible)}, not eligible: {len(opp_not)}")
    for _, row in opp_eligible.iterrows():
        assert row["meets_target"], f"{row['clinician_name']} eligible but doesn't meet target"
        assert row["utilization_pct"] >= 100, \
            f"{row['clinician_name']} eligible but util={row['utilization_pct']}"
    print(f"  PASS: OPP eligibility logic correct")

    # ------------------------------------------------------------------
    # Test 9: Excel workbook generation
    # ------------------------------------------------------------------
    print("\n[TEST 9] Excel workbook generation...")
    output_dir = Path("./output/test")
    output_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = output_dir / "QPi_Scorecard_AP_PP4_MOCK.xlsx"

    try:
        build_workbook(
            clinician_df=df,
            agency_metrics=am,
            pay_period=pp,
            clinic_key=clinic_key,
            output_path=xlsx_path,
        )
    except Exception as e:
        errors.append(f"Excel generation failed: {e}")
        print(f"  FAIL: {e}")
        import traceback
        traceback.print_exc()

    if xlsx_path.exists():
        size = xlsx_path.stat().st_size
        assert size > 5000, f"Excel file too small: {size} bytes"
        print(f"  PASS: Workbook generated ({size:,} bytes)")

        # Verify workbook structure
        import openpyxl
        wb = openpyxl.load_workbook(str(xlsx_path))
        expected_sheets = ["Clinician Scorecard", "Agency Roll-Up",
                           "Documentation Detail", "Cost Per Visit"]
        for sheet in expected_sheets:
            assert sheet in wb.sheetnames, f"Missing sheet: {sheet}"
        print(f"  PASS: All 4 tabs present: {wb.sheetnames}")

        # Verify clinician scorecard data rows
        ws1 = wb["Clinician Scorecard"]
        data_rows = ws1.max_row - 4  # subtract header rows
        assert data_rows == 15, f"Expected 15 data rows, got {data_rows}"
        print(f"  PASS: Clinician Scorecard has {data_rows} rows")

        # Verify agency roll-up has KPI rows
        ws2 = wb["Agency Roll-Up"]
        assert ws2["B7"].value == "Star Rating (SHP)", f"Row 7 KPI wrong: {ws2['B7'].value}"
        assert ws2["B15"].value == "Census", f"Row 15 KPI wrong: {ws2['B15'].value}"
        print(f"  PASS: Agency Roll-Up KPIs populated")

        wb.close()
    else:
        errors.append("Excel file not created")
        print(f"  FAIL: No file at {xlsx_path}")

    # ------------------------------------------------------------------
    # Test 10: Excel with trending data
    # ------------------------------------------------------------------
    print("\n[TEST 10] Excel workbook with trending data...")
    xlsx_trend_path = output_dir / "QPi_Scorecard_AP_PP4_MOCK_TREND.xlsx"
    try:
        build_workbook(
            clinician_df=df_trend,
            agency_metrics=am,
            pay_period=pp,
            clinic_key=clinic_key,
            output_path=xlsx_trend_path,
        )
        assert xlsx_trend_path.exists(), "Trend workbook not created"
        print(f"  PASS: Trend workbook generated ({xlsx_trend_path.stat().st_size:,} bytes)")
    except Exception as e:
        errors.append(f"Trend Excel generation failed: {e}")
        print(f"  FAIL: {e}")

    # ------------------------------------------------------------------
    # Test 11: CSV export
    # ------------------------------------------------------------------
    print("\n[TEST 11] CSV export...")
    csv_path = output_dir / "QPi_Raw_Data_PP4_MOCK.csv"
    df.to_csv(csv_path, index=False)
    assert csv_path.exists(), "CSV not created"
    csv_df = pd.read_csv(csv_path)
    assert len(csv_df) == 15, f"CSV row count wrong: {len(csv_df)}"
    print(f"  PASS: CSV exported ({len(csv_df)} rows, {len(csv_df.columns)} cols)")

    # ------------------------------------------------------------------
    # Test 12: Edge cases
    # ------------------------------------------------------------------
    print("\n[TEST 12] Edge cases...")

    # Empty productivity
    try:
        empty_df = build_clinician_scorecard(
            sf_productivity=pd.DataFrame(columns=mock_productivity().columns),
            sf_documentation=pd.DataFrame(),
            payroll=pd.DataFrame(),
        )
        assert len(empty_df) == 0, "Empty input should produce empty output"
        print(f"  PASS: Empty input handled gracefully")
    except Exception as e:
        errors.append(f"Empty input crashed: {e}")
        print(f"  FAIL: Empty input crashed: {e}")

    # Single clinician
    single = mock_productivity().head(1)
    single_doc = mock_documentation().head(1)
    try:
        single_df = build_clinician_scorecard(
            sf_productivity=single,
            sf_documentation=single_doc,
            payroll=pd.DataFrame(),
        )
        assert len(single_df) == 1, f"Single clinician: got {len(single_df)}"
        print(f"  PASS: Single clinician handled correctly")
    except Exception as e:
        errors.append(f"Single clinician crashed: {e}")
        print(f"  FAIL: Single clinician crashed: {e}")

    # ------------------------------------------------------------------
    # Test 13: Multi-agency
    # ------------------------------------------------------------------
    print("\n[TEST 13] Multi-agency workbook generation...")
    for ck, ag in AGENCIES.items():
        ag_path = output_dir / f"QPi_Scorecard_{ag.short_code}_PP4_MOCK.xlsx"
        try:
            build_workbook(
                clinician_df=df,
                agency_metrics=am,
                pay_period=pp,
                clinic_key=ck,
                output_path=ag_path,
            )
            assert ag_path.exists(), f"{ag.name} workbook not created"
            print(f"  PASS: {ag.name} ({ag.short_code}) workbook generated")
        except Exception as e:
            errors.append(f"{ag.name} workbook failed: {e}")
            print(f"  FAIL: {ag.name}: {e}")

    # ------------------------------------------------------------------
    # SUMMARY
    # ------------------------------------------------------------------
    print("\n" + "=" * 60)
    if errors:
        print(f"VALIDATION COMPLETE: {len(errors)} ERROR(S)")
        for e in errors:
            print(f"  ERROR: {e}")
        print("=" * 60)
        return False
    else:
        print("VALIDATION COMPLETE: ALL 13 TESTS PASSED")
        print(f"Output files in: {output_dir.resolve()}")
        print("=" * 60)
        return True


if __name__ == "__main__":
    success = run_test()
    sys.exit(0 if success else 1)
